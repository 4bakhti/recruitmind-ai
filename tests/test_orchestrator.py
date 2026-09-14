"""Tests for the LangGraph orchestrator.

Fully offline: every agent the nodes call is mocked, so these tests
exercise the wiring — which nodes run, in what order, and where the
conditional edges send a run — not the agents themselves.
"""

import unittest
from unittest.mock import patch

import orchestrator.nodes as nodes
from models.candidate import Candidate, GithubProfile, WorkEntry
from models.job_description import JobDescription, JobFitResult
from models.report import CandidateReport, ScoreBreakdown
from orchestrator import run_pipeline
from orchestrator.graph import (
    route_after_fuse,
    route_after_parse,
    route_after_score,
)
from orchestrator.state import initial_state


def make_candidate(**overrides) -> Candidate:
    data = {
        "name": "Ada Lovelace",
        "email": "ada@example.com",
        "skills": ["python", "fastapi"],
        "work_history": [
            WorkEntry(company="Acme", title="Engineer",
                      start_date="2020-01", end_date="Present")
        ],
        "raw_text": "Ada Lovelace — Engineer",
    }
    data.update(overrides)
    return Candidate(**data)


JOB = JobDescription(
    title="Backend Engineer",
    required_skills=["Python", "FastAPI"],
    description_text="Build REST APIs.",
)

FIT = JobFitResult(
    semantic_fit_score=78.0,
    hard_skill_match_score=100.0,
    composite_score=91.2,
    matched_required_skills=["Python", "FastAPI"],
)


def make_report(candidate: Candidate) -> CandidateReport:
    return CandidateReport(
        executive_summary="Strong backend candidate.",
        recommendation="Strong Fit",
        score_breakdown=ScoreBreakdown(
            composite_score=91.2, hard_skill_match_score=100.0,
            semantic_fit_score=78.0, hard_skill_weight=0.6,
            semantic_weight=0.4,
        ),
        candidate=candidate, job=JOB, job_fit=FIT,
    )


class RoutingTests(unittest.TestCase):
    """The conditional edges, tested as the plain functions they are."""

    def test_failed_parse_ends_the_run(self):
        state = initial_state("cv.pdf")
        self.assertEqual(route_after_parse(state), "end")

    def test_github_url_in_raw_text_routes_to_crawler(self):
        state = initial_state("cv.pdf")
        state["candidate"] = make_candidate(
            raw_text="Portfolio: https://github.com/ada"
        )
        self.assertEqual(route_after_parse(state), "crawl")

    def test_explicit_username_routes_to_crawler(self):
        state = initial_state("cv.pdf")
        state["candidate"] = make_candidate(github_username="ada")
        self.assertEqual(route_after_parse(state), "crawl")

    def test_no_github_signal_skips_to_fusion(self):
        state = initial_state("cv.pdf")
        state["candidate"] = make_candidate(raw_text="No links in this CV")
        self.assertEqual(route_after_parse(state), "fuse")

    def test_reserved_github_path_is_not_a_username(self):
        # github.com/features is a product page, not a profile
        state = initial_state("cv.pdf")
        state["candidate"] = make_candidate(
            raw_text="See https://github.com/features/copilot"
        )
        self.assertEqual(route_after_parse(state), "fuse")

    def test_fusion_ends_the_run_without_a_job(self):
        state = initial_state("cv.pdf")
        self.assertEqual(route_after_fuse(state), "end")

    def test_fusion_routes_to_scoring_with_a_job(self):
        state = initial_state("cv.pdf", JOB)
        self.assertEqual(route_after_fuse(state), "score")

    def test_no_scores_means_no_report(self):
        state = initial_state("cv.pdf", JOB)
        self.assertEqual(route_after_score(state), "end")

    def test_scores_route_to_the_report(self):
        state = initial_state("cv.pdf", JOB)
        state["fit"] = FIT
        self.assertEqual(route_after_score(state), "report")


class PipelineRunTests(unittest.TestCase):
    """End-to-end runs through the compiled graph, agents mocked."""

    def setUp(self):
        # Keep the retry test instant
        backoff = nodes.RETRY_BACKOFF_SECONDS
        nodes.RETRY_BACKOFF_SECONDS = 0
        self.addCleanup(setattr, nodes, "RETRY_BACKOFF_SECONDS", backoff)

    def test_full_run_with_github_and_job(self):
        candidate = make_candidate(raw_text="https://github.com/ada")
        enriched = candidate.model_copy(
            update={
                "github_username": "ada",
                "github_profile": GithubProfile(
                    public_repos=12, top_languages=["Go", "Python"]
                ),
            }
        )
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(nodes, "enrich_candidate", return_value=enriched), \
             patch.object(nodes, "score_fit", return_value=FIT), \
             patch.object(
                 nodes, "generate_report", return_value=make_report(enriched)
             ):
            state = run_pipeline("cv.pdf", JOB)

        self.assertEqual(
            state["steps_completed"],
            ["parse", "crawl", "fuse", "score", "report"],
        )
        self.assertEqual(state["errors"], [])
        self.assertIsNotNone(state["report"])
        # Fusion ran: skills normalized, GitHub-only language surfaced
        self.assertEqual(state["candidate"].skills, ["Python", "FastAPI"])
        self.assertEqual(state["candidate"].skills_from_github, ["Go"])

    def test_crawler_is_skipped_without_a_github_signal(self):
        candidate = make_candidate(raw_text="No links here")
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(nodes, "enrich_candidate") as crawler:
            state = run_pipeline("cv.pdf")

        crawler.assert_not_called()
        self.assertEqual(state["steps_completed"], ["parse", "fuse"])
        self.assertIn(
            "No GitHub data available", state["candidate"].data_quality_notes
        )

    def test_run_without_a_job_stops_after_fusion(self):
        candidate = make_candidate(raw_text="No links here")
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(nodes, "score_fit") as scorer, \
             patch.object(nodes, "generate_report") as reporter:
            state = run_pipeline("cv.pdf")

        scorer.assert_not_called()
        reporter.assert_not_called()
        self.assertIsNone(state["fit"])
        self.assertIsNone(state["report"])
        self.assertIsNotNone(state["candidate"])

    def test_parse_failure_is_retried_then_reported(self):
        with patch.object(nodes, "parse_cv", return_value=None) as parser:
            state = run_pipeline("cv.pdf", JOB)

        self.assertEqual(parser.call_count, nodes.PARSE_MAX_ATTEMPTS)
        self.assertIsNone(state["candidate"])
        self.assertEqual(state["steps_completed"], [])
        self.assertEqual(len(state["errors"]), 1)
        self.assertIn("Could not parse the CV", state["errors"][0])

    def test_parse_succeeds_on_retry(self):
        candidate = make_candidate(raw_text="No links here")
        with patch.object(
            nodes, "parse_cv", side_effect=[None, candidate]
        ) as parser:
            state = run_pipeline("cv.pdf")

        self.assertEqual(parser.call_count, 2)
        self.assertIsNotNone(state["candidate"])
        self.assertEqual(state["steps_completed"], ["parse", "fuse"])

    def test_failed_github_fetch_is_noted_not_fatal(self):
        candidate = make_candidate(raw_text="https://github.com/ada")
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(nodes, "enrich_candidate", return_value=candidate):
            state = run_pipeline("cv.pdf")

        # Crawl ran but produced nothing -> recorded, pipeline continued
        self.assertNotIn("crawl", state["steps_completed"])
        self.assertEqual(len(state["errors"]), 1)
        self.assertIn("could not fetch the profile", state["errors"][0])
        self.assertIsNotNone(state["candidate"])

    def test_scoring_failure_skips_the_report(self):
        candidate = make_candidate(raw_text="No links here")
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(
                 nodes, "score_fit", side_effect=RuntimeError("model missing")
             ), \
             patch.object(nodes, "generate_report") as reporter:
            state = run_pipeline("cv.pdf", JOB)

        reporter.assert_not_called()
        self.assertIsNone(state["fit"])
        self.assertIsNone(state["report"])
        self.assertIn("model missing", state["errors"][0])
        # The candidate we did manage to build is still returned
        self.assertIsNotNone(state["candidate"])

    def test_report_failure_leaves_the_scores_intact(self):
        candidate = make_candidate(raw_text="No links here")
        with patch.object(nodes, "parse_cv", return_value=candidate), \
             patch.object(nodes, "score_fit", return_value=FIT), \
             patch.object(
                 nodes, "generate_report",
                 side_effect=RuntimeError("no api key")
             ):
            state = run_pipeline("cv.pdf", JOB)

        self.assertEqual(state["fit"], FIT)
        self.assertIsNone(state["report"])
        self.assertIn("no api key", state["errors"][0])
        self.assertEqual(state["steps_completed"], ["parse", "fuse", "score"])

    def test_runs_do_not_leak_state_into_each_other(self):
        """The compiled graph is shared across calls; the state must not be."""
        candidate = make_candidate(raw_text="No links here")
        with patch.object(nodes, "parse_cv", return_value=candidate):
            first = run_pipeline("one.pdf")
            second = run_pipeline("two.pdf")

        self.assertEqual(first["file_path"], "one.pdf")
        self.assertEqual(second["file_path"], "two.pdf")
        self.assertEqual(second["steps_completed"], ["parse", "fuse"])
        self.assertEqual(second["errors"], [])


class RunPipelineEndpointTests(unittest.TestCase):
    """The /run-pipeline endpoint, driven through the ASGI app.

    No job description, so no embedding model is loaded. starlette 0.36
    (pinned via fastapi 0.109.2) is incompatible with the pinned httpx
    0.28 TestClient, so we drive the app through ASGITransport instead.
    """

    def test_trace_is_returned_for_a_run_without_a_job(self):
        import asyncio
        import tempfile
        from pathlib import Path

        import httpx
        from docx import Document

        import main

        candidate = make_candidate(raw_text="No links here")

        async def _post(cv_path):
            transport = httpx.ASGITransport(app=main.app)
            async with httpx.AsyncClient(
                transport=transport, base_url="http://testserver"
            ) as client:
                with open(cv_path, "rb") as fh:
                    return await client.post(
                        "/run-pipeline",
                        files={
                            "file": (
                                "cv.docx", fh,
                                "application/vnd.openxmlformats-"
                                "officedocument.wordprocessingml.document",
                            )
                        },
                    )

        with tempfile.TemporaryDirectory() as tmp:
            cv_path = Path(tmp) / "cv.docx"
            doc = Document()
            doc.add_paragraph("Ada Lovelace — Engineer")
            doc.save(str(cv_path))

            with patch.object(nodes, "parse_cv", return_value=candidate):
                response = asyncio.run(_post(cv_path))

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["candidate"]["name"], "Ada Lovelace")
        self.assertEqual(body["steps_completed"], ["parse", "fuse"])
        self.assertEqual(body["errors"], [])
        # No job supplied -> both stages skipped, reported as null
        self.assertIsNone(body["job_fit"])
        self.assertIsNone(body["report"])

    def test_unparseable_cv_is_a_422_carrying_the_reason(self):
        import asyncio
        import tempfile
        from pathlib import Path

        import httpx
        from docx import Document

        import main

        async def _post(cv_path):
            transport = httpx.ASGITransport(app=main.app)
            async with httpx.AsyncClient(
                transport=transport, base_url="http://testserver"
            ) as client:
                with open(cv_path, "rb") as fh:
                    return await client.post(
                        "/run-pipeline",
                        files={
                            "file": (
                                "cv.docx", fh,
                                "application/vnd.openxmlformats-"
                                "officedocument.wordprocessingml.document",
                            )
                        },
                    )

        backoff = nodes.RETRY_BACKOFF_SECONDS
        nodes.RETRY_BACKOFF_SECONDS = 0
        self.addCleanup(setattr, nodes, "RETRY_BACKOFF_SECONDS", backoff)

        with tempfile.TemporaryDirectory() as tmp:
            cv_path = Path(tmp) / "cv.docx"
            doc = Document()
            doc.add_paragraph("unreadable")
            doc.save(str(cv_path))

            with patch.object(nodes, "parse_cv", return_value=None):
                response = asyncio.run(_post(cv_path))

        self.assertEqual(response.status_code, 422)
        self.assertIn("Could not parse the CV", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
