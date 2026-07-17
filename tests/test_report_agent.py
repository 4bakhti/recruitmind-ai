"""Tests for agents/report_agent.py and the /full-report endpoint.

The Claude call is mocked everywhere — no API key needed. The end-to-end
test runs real local embeddings (model already cached by the job-fit
tests) but mocks Claude and GitHub.

Run from the project root:  python -m unittest tests.test_report_agent
"""

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from agents import report_agent
from agents.report_agent import (
    POSSIBLE_FIT_THRESHOLD,
    STRONG_FIT_THRESHOLD,
    recommendation_from_score,
    render_report_markdown,
    select_top_risks,
    select_top_strengths,
)
from models.candidate import (
    Candidate,
    GithubProfile,
    NotableRepo,
    WorkEntry,
)
from models.job_description import JobDescription, JobFitResult


def _fit(**overrides) -> JobFitResult:
    defaults = dict(
        semantic_fit_score=60.0,
        hard_skill_match_score=75.0,
        composite_score=69.0,
        matched_required_skills=["Python", "FastAPI"],
        gap_analysis=[],
        nice_to_have_matches=[],
        scoring_explanation="composite 69.0 = ...",
    )
    defaults.update(overrides)
    return JobFitResult(**defaults)


def _fake_claude_response(text: str):
    return SimpleNamespace(content=[SimpleNamespace(type="text", text=text)])


class RecommendationThresholdTests(unittest.TestCase):
    def test_strong_fit_band(self):
        self.assertEqual(recommendation_from_score(90.0), "Strong Fit")
        self.assertEqual(
            recommendation_from_score(STRONG_FIT_THRESHOLD), "Strong Fit"
        )

    def test_possible_fit_band(self):
        self.assertEqual(
            recommendation_from_score(STRONG_FIT_THRESHOLD - 0.1),
            "Possible Fit",
        )
        self.assertEqual(
            recommendation_from_score(POSSIBLE_FIT_THRESHOLD), "Possible Fit"
        )

    def test_weak_fit_band(self):
        self.assertEqual(
            recommendation_from_score(POSSIBLE_FIT_THRESHOLD - 0.1),
            "Weak Fit",
        )
        self.assertEqual(recommendation_from_score(0.0), "Weak Fit")


class StrengthSelectionTests(unittest.TestCase):
    def test_obvious_strengths_picked_in_priority_order(self):
        candidate = Candidate(
            name="Star Dev",
            github_profile=GithubProfile(
                public_repos=20,
                total_stars=500,
                notable_repos=[
                    NotableRepo(name="bigproject", stars=400, language="Python")
                ],
                contribution_summary="Active — 15 of 20 public repos pushed "
                                     "to in the last year",
            ),
        )
        fit = _fit(nice_to_have_matches=["Kubernetes"])
        strengths = select_top_strengths(candidate, fit)

        self.assertEqual(len(strengths), 3)  # capped at 3
        self.assertIn("Matches 2/2 required skills", strengths[0])
        self.assertIn("Kubernetes", strengths[1])
        self.assertIn("bigproject", strengths[2])

    def test_work_history_demonstrating_matched_skill(self):
        candidate = Candidate(
            work_history=[
                WorkEntry(
                    company="Acme",
                    title="Backend Engineer",
                    description="Built services with Python and FastAPI.",
                )
            ],
        )
        strengths = select_top_strengths(candidate, _fit())
        self.assertTrue(any("Acme" in s for s in strengths))

    def test_no_signals_means_no_strengths(self):
        fit = _fit(matched_required_skills=[], hard_skill_match_score=0.0)
        self.assertEqual(select_top_strengths(Candidate(), fit), [])


class RiskSelectionTests(unittest.TestCase):
    def test_obvious_risks_picked_in_priority_order(self):
        candidate = Candidate(
            employment_gap_notes=[
                "Employment gap of ~9 months between Acme (ended 2020-06) "
                "and Beta (started 2021-03)"
            ],
            data_quality_notes=["No GitHub data available"],
        )
        fit = _fit(
            gap_analysis=["Kubernetes", "Terraform"],
            hard_skill_match_score=33.3,
            semantic_fit_score=25.0,
        )
        risks = select_top_risks(candidate, fit)

        self.assertEqual(len(risks), 3)  # capped at 3
        self.assertIn("Missing required skills: Kubernetes, Terraform", risks[0])
        self.assertIn("hard-skill score 33.3/100", risks[1])
        self.assertIn("Employment gap", risks[2])

    def test_clean_candidate_has_no_risks(self):
        risks = select_top_risks(Candidate(), _fit())
        self.assertEqual(risks, [])


class MarkdownRenderTests(unittest.TestCase):
    def test_report_contains_all_sections_and_key_values(self):
        candidate = Candidate(name="Jane Doe", skills=["Python"])
        job = JobDescription(
            title="Backend Engineer",
            company="RecruitCo",
            required_skills=["Python", "FastAPI"],
        )
        fit = _fit(composite_score=81.0, gap_analysis=["FastAPI"])

        with patch.object(report_agent, "_get_client") as get_client:
            get_client.return_value.messages.create.return_value = (
                _fake_claude_response("Jane Doe is a strong backend hire.")
            )
            report = report_agent.generate_report(candidate, job, fit)

        md = report.report_markdown
        self.assertEqual(md, render_report_markdown(report))
        self.assertIn("# Candidate Report: Jane Doe", md)
        self.assertIn("Backend Engineer — RecruitCo", md)
        self.assertIn("**Recommendation:** Strong Fit", md)
        self.assertIn("## Executive Summary", md)
        self.assertIn("Jane Doe is a strong backend hire.", md)
        self.assertIn("## Top Strengths", md)
        self.assertIn("## Top Risks", md)
        self.assertIn("## Score Breakdown", md)
        self.assertIn("| Hard skill match | 75.0 | 60% |", md)
        self.assertIn("| Semantic fit | 60.0 | 40% |", md)
        self.assertIn("**81.0**", md)
        self.assertIn("Missing required skills: FastAPI", md)

    def test_llm_failure_falls_back_to_deterministic_summary(self):
        candidate = Candidate(name="Jane Doe")
        job = JobDescription(title="Backend Engineer")
        with patch.object(report_agent, "_get_client") as get_client:
            get_client.return_value.messages.create.side_effect = (
                RuntimeError("API down")
            )
            report = report_agent.generate_report(candidate, job, _fit())

        self.assertIn("Jane Doe", report.executive_summary)
        self.assertIn("69.0/100", report.executive_summary)
        # The deterministic parts are unaffected by the LLM failure
        self.assertEqual(report.recommendation, "Possible Fit")


CLAUDE_PARSE_RESPONSE = json.dumps(
    {
        "name": "Jane Doe",
        "email": "jane.doe@example.com",
        "phone": None,
        "location": "Berlin, Germany",
        "education": [],
        "work_history": [
            {
                "company": "Acme GmbH",
                "title": "Backend Engineer",
                "start_date": "2020-01",
                "end_date": "Present",
                "description": "Built REST APIs with Python and FastAPI.",
            }
        ],
        "skills": ["Python", "FastAPI", "PostgreSQL"],
        "raw_text": None,
    }
)

JOB_DESCRIPTION_JSON = json.dumps(
    {
        "title": "Backend Engineer",
        "company": "RecruitCo",
        "required_skills": ["Python", "FastAPI"],
        "nice_to_have_skills": ["Kubernetes"],
        "description_text": "Build and operate REST APIs in Python.",
    }
)


class FullReportEndpointTests(unittest.TestCase):
    # starlette 0.36 (pinned via fastapi 0.109.2) TestClient is incompatible
    # with the pinned httpx 0.28, so we drive the ASGI app through
    # httpx.ASGITransport instead — same end-to-end path.
    @staticmethod
    def _post_full_report(app, cv_path: Path):
        import asyncio
        import httpx

        async def _post():
            transport = httpx.ASGITransport(app=app)
            async with httpx.AsyncClient(
                transport=transport, base_url="http://testserver"
            ) as client:
                with open(cv_path, "rb") as fh:
                    return await client.post(
                        "/full-report",
                        files={
                            "file": (
                                "cv.docx",
                                fh,
                                "application/vnd.openxmlformats-"
                                "officedocument.wordprocessingml.document",
                            )
                        },
                        data={"job_description": JOB_DESCRIPTION_JSON},
                    )

        return asyncio.run(_post())

    def test_full_report_end_to_end(self):
        import main

        with tempfile.TemporaryDirectory() as tmp:
            cv_path = Path(tmp) / "cv.docx"
            doc = Document()
            doc.add_paragraph("Jane Doe — Backend Engineer")
            doc.add_paragraph("Python, FastAPI, PostgreSQL")
            doc.save(str(cv_path))

            with patch("agents.parser_agent._get_client") as parse_client, \
                 patch.object(report_agent, "_get_client") as report_client, \
                 patch(
                     "agents.github_crawler_agent.fetch_github_profile",
                     return_value=None,
                 ):
                parse_client.return_value.messages.create.return_value = (
                    _fake_claude_response(CLAUDE_PARSE_RESPONSE)
                )
                report_client.return_value.messages.create.return_value = (
                    _fake_claude_response(
                        "Jane Doe is a strong match. Hire her."
                    )
                )

                response = self._post_full_report(main.app, cv_path)

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()

        self.assertIn(
            body["recommendation"], ["Strong Fit", "Possible Fit", "Weak Fit"]
        )
        self.assertEqual(
            body["executive_summary"], "Jane Doe is a strong match. Hire her."
        )
        # Both required skills present -> full hard-skill score, no gaps
        self.assertEqual(body["job_fit"]["hard_skill_match_score"], 100.0)
        self.assertEqual(body["job_fit"]["gap_analysis"], [])
        self.assertEqual(body["candidate"]["name"], "Jane Doe")
        self.assertEqual(body["job"]["title"], "Backend Engineer")
        self.assertIn("# Candidate Report: Jane Doe", body["report_markdown"])
        self.assertIn("## Score Breakdown", body["report_markdown"])
        # GitHub crawl was mocked to fail -> noted as a data-quality risk
        self.assertIn(
            "No GitHub data available",
            body["candidate"]["data_quality_notes"],
        )


if __name__ == "__main__":
    unittest.main()
